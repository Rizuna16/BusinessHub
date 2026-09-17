import csv
import io
from datetime import datetime, date
from decimal import Decimal
from enum import Enum
from typing import Any, Dict, List, Optional



def _sanitize_formula_injection(value: str) -> str:
    if not isinstance(value, str):
        return value
    stripped = value.lstrip()
    if stripped and stripped[0] in ("=", "+", "-", "@"):
        return "\u200b" + value
    return value


def _serialize_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    # Decimal, int, float are always numeric — never sanitize (preserve negative numbers)
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (int, float)):
        return str(value)
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, Enum):
        return value.value if hasattr(value, "value") else str(value)
    # Only strings are eligible for formula-injection sanitization
    return _sanitize_formula_injection(str(value))
    if not isinstance(value, str):
        return value
    stripped = value.lstrip()
    if stripped and stripped[0] in ("=", "+", "-", "@"):
        return "\u200b" + value
    return value


def _serialize_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (int, float)):
        return str(value)
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, Enum):
        return value.value if hasattr(value, "value") else str(value)
    return _sanitize_formula_injection(str(value))


def generate_csv(
    headers: List[str],
    rows: List[Dict[str, Any]],
) -> bytes:
    output = io.StringIO()
    output.write("\ufeff")
    writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)
    writer.writerow(headers)
    for row in rows:
        writer.writerow([_serialize_value(row.get(h)) for h in headers])
    return output.getvalue().encode("utf-8")


def generate_xlsx(
    headers: List[str],
    rows: List[Dict[str, Any]],
    sheet_name: str = "Export",
) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    for row_idx, row_data in enumerate(rows, 2):
        for col_idx, header in enumerate(headers, 1):
            value = row_data.get(header)
            cell = ws.cell(row=row_idx, column=col_idx)
            if value is None:
                cell.value = ""
            elif isinstance(value, Decimal):
                cell.value = str(value)
                cell.number_format = "@"
            elif isinstance(value, (int,)) and not isinstance(value, bool):
                cell.value = value
            elif isinstance(value, bool):
                cell.value = "Yes" if value else "No"
            elif isinstance(value, datetime):
                cell.value = value.strftime("%Y-%m-%d %H:%M:%S")
            elif isinstance(value, date):
                cell.value = value.strftime("%Y-%m-%d")
            else:
                cell.value = _sanitize_formula_injection(str(value))
            cell.border = thin_border

    for col_idx, header in enumerate(headers, 1):
        max_length = len(str(header))
        for row_idx in range(2, len(rows) + 2):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val is not None:
                max_length = max(max_length, len(str(val)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_length + 4, 50)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()
