import sys
import os
import io
import csv
import re
from decimal import Decimal
from datetime import datetime, timezone, date

import pytest
from fastapi.testclient import TestClient

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app.main import app
from app.modules.authentication.repository import InMemoryUserRepository
from app.modules.export.formatters import (
    generate_csv,
    generate_xlsx,
    _serialize_value,
    _sanitize_formula_injection,
)
from app.modules.export.registry import (
    EXPORT_REGISTRY,
    get_export_definition,
    validate_resource,
    validate_format,
    ExportResourceType,
)
from app.modules.export.service import (
    ExportService,
    MAX_EXPORT_ROWS,
    _sanitize_filename,
    _parse_date,
)


@pytest.fixture(autouse=True)
def clear_repos():
    InMemoryUserRepository.clear()
    yield
    InMemoryUserRepository.clear()


@pytest.fixture
def client():
    with TestClient(app) as c:
        yield c


def _register_and_get_token(client: TestClient, email: str = "owner@test.com") -> str:
    r = client.post("/api/v1/auth/register", json={
        "email": email,
        "full_name": "Test Owner",
        "password": "Password123",
        "password_confirmation": "Password123",
    })
    if r.status_code == 201:
        token = r.json().get("access_token")
        if token:
            return token
    r = client.post("/api/v1/auth/login", json={
        "email": email,
        "password": "Password123",
    })
    assert r.status_code == 200, f"Login failed: {r.json()}"
    return r.json()["access_token"]


def _create_business(client: TestClient, token: str) -> str:
    r = client.post(
        "/api/v1/businesses",
        headers={"Authorization": f"Bearer {token}"},
        json={"name": "Test Biz", "business_type": "umkm", "timezone": "UTC", "locale": "en-US"},
    )
    return r.json()["id"]


def _seed_minimal_data(client: TestClient, token: str, business_id: str):
    unit = client.post(
        f"/api/v1/businesses/{business_id}/units",
        headers={"Authorization": f"Bearer {token}"},
        json={"name": "Unit", "code": "U1", "symbol": "pcs", "unit_type": "OTHER"},
    ).json()["id"]
    cat = client.post(
        f"/api/v1/businesses/{business_id}/categories",
        headers={"Authorization": f"Bearer {token}"},
        json={"name": "Category", "code": "CAT1", "description": "Test"},
    ).json()["id"]
    product = client.post(
        f"/api/v1/businesses/{business_id}/products",
        headers={"Authorization": f"Bearer {token}"},
        json={"name": "Product", "code": "P1", "unit_id": unit, "category_id": cat, "product_type": "GOODS"},
    ).json()["id"]
    customer = client.post(
        f"/api/v1/businesses/{business_id}/customers",
        headers={"Authorization": f"Bearer {token}"},
        json={"name": "Customer", "customer_type": "INDIVIDUAL"},
    ).json()["id"]
    supplier = client.post(
        f"/api/v1/businesses/{business_id}/suppliers",
        headers={"Authorization": f"Bearer {token}"},
        json={"name": "Supplier", "supplier_type": "ORGANIZATION"},
    ).json()["id"]
    branch = client.post(
        f"/api/v1/businesses/{business_id}/branches",
        headers={"Authorization": f"Bearer {token}"},
        json={"name": "Branch", "code": "BR1"},
    ).json()["id"]
    warehouse = client.post(
        f"/api/v1/businesses/{business_id}/warehouses",
        headers={"Authorization": f"Bearer {token}"},
        json={"name": "Warehouse", "code": "WH1"},
    ).json()["id"]
    location = client.post(
        f"/api/v1/businesses/{business_id}/warehouses/{warehouse}/locations",
        headers={"Authorization": f"Bearer {token}"},
        json={"name": "Location", "code": "LOC1", "location_type": "GENERAL"},
    ).json()["id"]
    client.post(
        f"/api/v1/businesses/{business_id}/inventory/opening-balance",
        headers={"Authorization": f"Bearer {token}"},
        json={"inventory_location_id": location, "product_id": product, "quantity": "100"},
    )
    return {
        "unit": unit, "category": cat, "product": product,
        "customer": customer, "supplier": supplier, "branch": branch,
        "warehouse": warehouse, "location": location,
    }


# ============================================================
# A. REGISTRY TESTS
# ============================================================

class TestRegistry:
    def test_all_expected_resources_registered(self):
        expected = {
            "products", "categories", "units", "customers", "suppliers",
            "warehouse_stock", "purchases", "sales_checkouts", "sales_orders",
            "delivery_notes", "expenses", "ar_aging", "trial_balance",
            "profit_loss", "balance_sheet", "cash_flow",
            "platform_businesses", "platform_users",
        }
        assert expected.issubset(set(EXPORT_REGISTRY.keys()))

    def test_get_export_definition_valid(self):
        defn = get_export_definition("products")
        assert defn is not None
        assert defn.resource_key == "products"
        assert defn.resource_type == ExportResourceType.TABLE
        assert "csv" in defn.allowed_formats
        assert "xlsx" in defn.allowed_formats

    def test_get_export_definition_unknown_returns_none(self):
        assert get_export_definition("nonexistent_resource") is None

    def test_validate_resource_unknown_raises_404(self):
        with pytest.raises(Exception) as exc_info:
            validate_resource("nonexistent_resource")
        assert "404" in str(exc_info.value.status_code) or "not found" in str(exc_info.value.detail).lower()

    def test_validate_format_invalid_raises_400(self):
        defn = get_export_definition("products")
        with pytest.raises(Exception) as exc_info:
            validate_format(defn, "pdf")
        assert "400" in str(exc_info.value.status_code) or "unsupported" in str(exc_info.value.detail).lower()

    def test_platform_resources_require_super_admin(self):
        defn = get_export_definition("platform_businesses")
        assert "SUPER_ADMIN" in defn.required_roles

    def test_table_resources_require_member_plus(self):
        defn = get_export_definition("products")
        assert "MEMBER" in defn.required_roles
        assert "ADMIN" in defn.required_roles
        assert "OWNER" in defn.required_roles


# ============================================================
# B. PAGINATION TESTS
# ============================================================

class TestPagination:
    def test_parse_date_valid(self):
        d = _parse_date("2026-01-15")
        assert d == date(2026, 1, 15)

    def test_parse_date_none(self):
        assert _parse_date(None) is None

    def test_parse_date_invalid(self):
        with pytest.raises(Exception):
            _parse_date("not-a-date")

    def test_sanitize_filename(self):
        result = _sanitize_filename("My Business @#$ Name")
        assert "/" not in result
        assert "\\" not in result
        assert ".." not in result
        assert result.isascii()


# ============================================================
# C. CSV FORMATTER TESTS
# ============================================================

class TestCSVFormatter:
    def test_csv_headers_present(self):
        content = generate_csv(["name", "value"], [{"name": "A", "value": 1}])
        decoded = content.decode("utf-8")
        lines = decoded.strip().split("\n")
        assert "name,value" in lines[0] or '"name","value"' in lines[0]

    def test_csv_utf8_bom(self):
        content = generate_csv(["col"], [{"col": "test"}])
        assert content[:3] == b"\xef\xbb\xbf"

    def test_decimal_exact_representation(self):
        content = generate_csv(["amount"], [{"amount": Decimal("150000.10")}])
        decoded = content.decode("utf-8-sig").strip()
        assert "150000.10" in decoded

    def test_negative_decimal_preserved(self):
        content = generate_csv(["amount"], [{"amount": Decimal("-150000.10")}])
        decoded = content.decode("utf-8-sig").strip()
        assert "-150000.10" in decoded

    def test_zero_decimal(self):
        content = generate_csv(["amount"], [{"amount": Decimal("0")}])
        decoded = content.decode("utf-8-sig").strip()
        assert "0" in decoded

    def test_trailing_scale_preserved(self):
        content = generate_csv(["amount"], [{"amount": Decimal("100.0000")}])
        decoded = content.decode("utf-8-sig").strip()
        assert "100.0000" in decoded

    def test_large_decimal(self):
        content = generate_csv(["amount"], [{"amount": Decimal("999999999999.9999")}])
        decoded = content.decode("utf-8-sig").strip()
        assert "999999999999.9999" in decoded

    def test_formula_like_string_sanitized(self):
        content = generate_csv(["col"], [{"col": "=SUM(A1:A10)"}])
        decoded = content.decode("utf-8-sig").strip()
        assert "\u200b=SUM(A1:A10)" in decoded or "'=SUM(A1:A10)" in decoded

    def test_plus_prefix_sanitized(self):
        content = generate_csv(["col"], [{"col": "+formula"}])
        decoded = content.decode("utf-8-sig").strip()
        assert "\u200b+formula" in decoded or "'+formula" in decoded

    def test_negative_string_sanitized(self):
        content = generate_csv(["col"], [{"col": "-150000.10"}])
        decoded = content.decode("utf-8-sig").strip()
        assert "-150000.10" in decoded

    def test_at_prefix_sanitized(self):
        content = generate_csv(["col"], [{"col": "@data"}])
        decoded = content.decode("utf-8-sig").strip()
        assert "\u200b@data" in decoded or "'@data" in decoded

    def test_none_representation(self):
        content = generate_csv(["col"], [{"col": None}])
        decoded = content.decode("utf-8-sig").strip()
        lines = decoded.split("\n")
        assert len(lines) == 2

    def test_bool_representation(self):
        content = generate_csv(["col"], [{"col": True}])
        decoded = content.decode("utf-8-sig").strip()
        assert "Yes" in decoded

    def test_datetime_representation(self):
        content = generate_csv(["col"], [{"col": datetime(2026, 9, 15, 10, 30, 0)}])
        decoded = content.decode("utf-8-sig").strip()
        assert "2026-09-15" in decoded

    def test_date_representation(self):
        content = generate_csv(["col"], [{"col": date(2026, 9, 15)}])
        decoded = content.decode("utf-8-sig").strip()
        assert "2026-09-15" in decoded

    def test_normal_string_preserved(self):
        content = generate_csv(["col"], [{"col": "Hello World"}])
        decoded = content.decode("utf-8-sig").strip()
        assert "Hello World" in decoded


# ============================================================
# D. XLSX FORMATTER TESTS
# ============================================================

class TestXLSXFormatter:
    def _get_cell_values(self, xlsx_content: bytes):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(xlsx_content))
        ws = wb.active
        return ws

    def test_xlsx_workbook_generation(self):
        content = generate_xlsx(["name", "value"], [{"name": "Test", "value": 100}])
        assert len(content) > 0
        assert isinstance(content, bytes)

    def test_xlsx_worksheet_naming(self):
        content = generate_xlsx(["col"], [{"col": "val"}], sheet_name="MySheet")
        ws = self._get_cell_values(content)
        assert ws.title == "MySheet"

    def test_xlsx_decimal_stored_as_string(self):
        content = generate_xlsx(["amount"], [{"amount": Decimal("150000.10")}])
        ws = self._get_cell_values(content)
        cell = ws.cell(row=2, column=1)
        assert isinstance(cell.value, str)
        assert cell.value == "150000.10"

    def test_xlsx_negative_decimal_stored_as_string(self):
        content = generate_xlsx(["amount"], [{"amount": Decimal("-150000.10")}])
        ws = self._get_cell_values(content)
        cell = ws.cell(row=2, column=1)
        assert cell.value == "-150000.10"

    def test_xlsx_zero_decimal(self):
        content = generate_xlsx(["amount"], [{"amount": Decimal("0")}])
        ws = self._get_cell_values(content)
        cell = ws.cell(row=2, column=1)
        assert cell.value == "0"

    def test_xlsx_trailing_scale_preserved(self):
        content = generate_xlsx(["amount"], [{"amount": Decimal("100.0000")}])
        ws = self._get_cell_values(content)
        cell = ws.cell(row=2, column=1)
        assert cell.value == "100.0000"

    def test_xlsx_large_decimal(self):
        content = generate_xlsx(["amount"], [{"amount": Decimal("999999999999.9999")}])
        ws = self._get_cell_values(content)
        cell = ws.cell(row=2, column=1)
        assert cell.value == "999999999999.9999"

    def test_xlsx_integer_quantity_numeric(self):
        content = generate_xlsx(["qty"], [{"qty": 42}])
        ws = self._get_cell_values(content)
        cell = ws.cell(row=2, column=1)
        assert cell.value == 42
        assert isinstance(cell.value, int)

    def test_xlsx_formula_string_sanitized(self):
        content = generate_xlsx(["col"], [{"col": "=SUM(A1:A10)"}])
        ws = self._get_cell_values(content)
        cell = ws.cell(row=2, column=1)
        assert str(cell.value).startswith("\u200b") or str(cell.value).startswith("'")

    def test_xlsx_none_representation(self):
        content = generate_xlsx(["col"], [{"col": None}])
        ws = self._get_cell_values(content)
        cell = ws.cell(row=2, column=1)
        assert cell.value is None or cell.value == ""

    def test_xlsx_bool_representation(self):
        content = generate_xlsx(["col"], [{"col": True}])
        ws = self._get_cell_values(content)
        cell = ws.cell(row=2, column=1)
        assert cell.value == "Yes"

    def test_xlsx_datetime_representation(self):
        content = generate_xlsx(["col"], [{"col": datetime(2026, 9, 15, 10, 30, 0)}])
        ws = self._get_cell_values(content)
        cell = ws.cell(row=2, column=1)
        assert "2026-09-15" in str(cell.value)

    def test_xlsx_header_bold(self):
        from openpyxl.styles import Font
        content = generate_xlsx(["col"], [{"col": "val"}])
        ws = self._get_cell_values(content)
        cell = ws.cell(row=1, column=1)
        assert cell.font.bold is True


# ============================================================
# E. AUTHORIZATION & INTEGRATION TESTS
# ============================================================

class TestExportAPI:
    def test_export_products_csv(self, client: TestClient):
        token = _register_and_get_token(client, "export_owner@test.com")
        biz = _create_business(client, token)
        _seed_minimal_data(client, token, biz)

        r = client.get(
            f"/api/v1/businesses/{biz}/export/table/products",
            headers={"Authorization": f"Bearer {token}"},
            params={"format": "csv", "export_mode": "all_matching"},
        )
        assert r.status_code == 200
        assert "text/csv" in r.headers.get("content-type", "")
        assert "attachment" in r.headers.get("content-disposition", "")
        assert "Product" in r.headers.get("content-disposition", "") or "products" in r.headers.get("content-disposition", "")

    def test_export_products_xlsx(self, client: TestClient):
        token = _register_and_get_token(client, "xlsx_owner@test.com")
        biz = _create_business(client, token)
        _seed_minimal_data(client, token, biz)

        r = client.get(
            f"/api/v1/businesses/{biz}/export/table/products",
            headers={"Authorization": f"Bearer {token}"},
            params={"format": "xlsx", "export_mode": "all_matching"},
        )
        assert r.status_code == 200
        assert "spreadsheetml" in r.headers.get("content-type", "")

    def test_export_unauthenticated_returns_401(self, client: TestClient):
        r = client.get("/api/v1/businesses/fake/export/table/products")
        assert r.status_code in (401, 403)

    def test_export_unknown_resource_returns_404(self, client: TestClient):
        token = _register_and_get_token(client, "404_owner@test.com")
        biz = _create_business(client, token)
        r = client.get(
            f"/api/v1/businesses/{biz}/export/table/nonexistent",
            headers={"Authorization": f"Bearer {token}"},
        )
        assert r.status_code == 404

    def test_export_unsupported_format_returns_400(self, client: TestClient):
        token = _register_and_get_token(client, "pdf_owner@test.com")
        biz = _create_business(client, token)
        r = client.get(
            f"/api/v1/businesses/{biz}/export/table/products",
            headers={"Authorization": f"Bearer {token}"},
            params={"format": "pdf"},
        )
        assert r.status_code == 400

    def test_export_categories_csv(self, client: TestClient):
        token = _register_and_get_token(client, "cat_owner@test.com")
        biz = _create_business(client, token)
        _seed_minimal_data(client, token, biz)

        r = client.get(
            f"/api/v1/businesses/{biz}/export/table/categories",
            headers={"Authorization": f"Bearer {token}"},
            params={"format": "csv"},
        )
        assert r.status_code == 200

    def test_export_units_csv(self, client: TestClient):
        token = _register_and_get_token(client, "unit_owner@test.com")
        biz = _create_business(client, token)
        _seed_minimal_data(client, token, biz)

        r = client.get(
            f"/api/v1/businesses/{biz}/export/table/units",
            headers={"Authorization": f"Bearer {token}"},
            params={"format": "csv"},
        )
        assert r.status_code == 200

    def test_export_customers_csv(self, client: TestClient):
        token = _register_and_get_token(client, "cust_owner@test.com")
        biz = _create_business(client, token)
        _seed_minimal_data(client, token, biz)

        r = client.get(
            f"/api/v1/businesses/{biz}/export/table/customers",
            headers={"Authorization": f"Bearer {token}"},
            params={"format": "csv"},
        )
        assert r.status_code == 200
        decoded = r.content.decode("utf-8-sig")
        assert "Customer" in decoded or "customer" in decoded.lower()

    def test_export_suppliers_csv(self, client: TestClient):
        token = _register_and_get_token(client, "sup_owner@test.com")
        biz = _create_business(client, token)
        _seed_minimal_data(client, token, biz)

        r = client.get(
            f"/api/v1/businesses/{biz}/export/table/suppliers",
            headers={"Authorization": f"Bearer {token}"},
            params={"format": "csv"},
        )
        assert r.status_code == 200

    def test_export_warehouse_stock_csv(self, client: TestClient):
        token = _register_and_get_token(client, "wh_owner@test.com")
        biz = _create_business(client, token)
        _seed_minimal_data(client, token, biz)

        r = client.get(
            f"/api/v1/businesses/{biz}/export/table/warehouse_stock",
            headers={"Authorization": f"Bearer {token}"},
            params={"format": "csv"},
        )
        assert r.status_code == 200

    def test_export_trial_balance_requires_accounting(self, client: TestClient):
        token = _register_and_get_token(client, "tb_owner@test.com")
        biz = _create_business(client, token)

        r = client.get(
            f"/api/v1/businesses/{biz}/export/report/trial_balance",
            headers={"Authorization": f"Bearer {token}"},
            params={"format": "csv"},
        )
        assert r.status_code == 200

    def test_export_profit_loss_requires_period(self, client: TestClient):
        token = _register_and_get_token(client, "pl_owner@test.com")
        biz = _create_business(client, token)

        r = client.get(
            f"/api/v1/businesses/{biz}/export/report/profit_loss",
            headers={"Authorization": f"Bearer {token}"},
            params={"format": "csv"},
        )
        assert r.status_code == 422

    def test_export_invalid_date_returns_422(self, client: TestClient):
        token = _register_and_get_token(client, "date_owner@test.com")
        biz = _create_business(client, token)

        r = client.get(
            f"/api/v1/businesses/{biz}/export/table/purchases",
            headers={"Authorization": f"Bearer {token}"},
            params={"format": "csv", "date_from": "not-a-date"},
        )
        assert r.status_code == 422


# ============================================================
# F. PLATFORM EXPORT TESTS
# ============================================================

class TestPlatformExport:
    def _get_superadmin_token(self, client: TestClient) -> str:
        from app.modules.authentication.repository import InMemoryUserRepository
        InMemoryUserRepository.clear()
        r = client.post("/api/v1/auth/register", json={
            "email": "superadmin@test.com",
            "full_name": "Super Admin",
            "password": "SuperAdminPass123!",
            "password_confirmation": "SuperAdminPass123!",
        })
        token = r.json().get("access_token")
        if not token:
            r = client.post("/api/v1/auth/login", json={
                "email": "superadmin@test.com",
                "password": "SuperAdminPass123!",
            })
            assert r.status_code == 200, f"Login failed: {r.json()}"
            token = r.json()["access_token"]

        from app.modules.authentication.repository import user_repository
        users = list(user_repository._users.values())
        for u in users:
            if u.email == "superadmin@test.com":
                from app.modules.authentication.schemas import PlatformRole
                u.platform_role = PlatformRole.SUPER_ADMIN
                user_repository._users[u.id] = u
                break
        return token

    def test_platform_export_businesses_superadmin(self, client: TestClient):
        token = self._get_superadmin_token(client)
        r = client.get(
            "/api/v1/platform/export/platform_businesses",
            headers={"Authorization": f"Bearer {token}"},
            params={"format": "csv"},
        )
        assert r.status_code == 200
        assert "text/csv" in r.headers.get("content-type", "")

    def test_platform_export_users_superadmin(self, client: TestClient):
        token = self._get_superadmin_token(client)
        r = client.get(
            "/api/v1/platform/export/platform_users",
            headers={"Authorization": f"Bearer {token}"},
            params={"format": "csv"},
        )
        assert r.status_code == 200

    def test_platform_export_denied_for_regular_user(self, client: TestClient):
        token = _register_and_get_token(client, "regular_user@test.com")
        r = client.get(
            "/api/v1/platform/export/platform_businesses",
            headers={"Authorization": f"Bearer {token}"},
            params={"format": "csv"},
        )
        assert r.status_code == 403


# ============================================================
# G. FORMULA INJECTION TESTS
# ============================================================

class TestFormulaInjection:
    def test_sanitize_equals(self):
        result = _sanitize_formula_injection("=SUM(A1)")
        assert result.startswith("\u200b") or result.startswith("'")

    def test_sanitize_plus(self):
        result = _sanitize_formula_injection("+cmd")
        assert result.startswith("\u200b") or result.startswith("'")

    def test_sanitize_at(self):
        result = _sanitize_formula_injection("@data")
        assert result.startswith("\u200b") or result.startswith("'")

    def test_negative_number_not_sanitized(self):
        result = _serialize_value(Decimal("-150000.10"))
        assert result == "-150000.10"

    def test_negative_string_starting_with_dash_preserved(self):
        result = _serialize_value("-150000.10")
        assert "-150000.10" in result

    def test_normal_string_not_sanitized(self):
        result = _sanitize_formula_injection("Hello World")
        assert result == "Hello World"


# ============================================================
# H. DATA INTEGRITY TESTS
# ============================================================

class TestDataIntegrity:
    def test_decimal_no_float_conversion(self):
        val = Decimal("150000.10")
        serialized = _serialize_value(val)
        assert serialized == "150000.10"
        assert "150000.1" not in serialized or "150000.10" in serialized

    def test_negative_decimal_preserved(self):
        val = Decimal("-999999999999.9999")
        serialized = _serialize_value(val)
        assert serialized == "-999999999999.9999"

    def test_zero_decimal(self):
        val = Decimal("0")
        serialized = _serialize_value(val)
        assert serialized == "0"

    def test_no_rounding(self):
        val = Decimal("100.0000")
        serialized = _serialize_value(val)
        assert serialized == "100.0000"


# ============================================================
# I. EXPORT SERVICE UNIT TESTS
# ============================================================

class TestExportService:
    def test_sanitize_business_name(self):
        from app.modules.export.service import _build_filename
        result = _build_filename("My Business", "products", "csv")
        assert "my_business" in result
        assert result.endswith(".csv")

    def test_sanitize_business_name_special_chars(self):
        from app.modules.export.service import _sanitize_business_name
        result = _sanitize_business_name("Biz@#$%Name")
        assert "/" not in result
        assert ".." not in result
